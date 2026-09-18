from __future__ import annotations

import io
import re
from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document as DocxDocument

from app.utils.error_handlers import APIError
from app.utils.text_utils import normalize_text


EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{8,}\d)")
SECTION_HEADERS = {
    "skills": re.compile(r"^(skills|technical skills|core competencies)\b", re.I),
    "education": re.compile(r"^(education|academic|qualifications)\b", re.I),
    "experience": re.compile(r"^(experience|work experience|employment|professional experience)\b", re.I),
    "projects": re.compile(r"^(projects|personal projects|academic projects)\b", re.I),
    "certifications": re.compile(r"^(certifications|certificates|licenses)\b", re.I),
}


def extract_text_from_pdf(source) -> str:
    try:
        reader = PdfReader(source)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return normalize_text("\n".join(pages))
    except Exception as exc:
        raise APIError("INVALID_DOCUMENT", f"Unable to read PDF: {exc}", 400) from exc


def extract_text_from_docx(source) -> str:
    try:
        document = DocxDocument(source)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return normalize_text("\n".join(parts))
    except Exception as exc:
        raise APIError("INVALID_DOCUMENT", f"Unable to read DOCX: {exc}", 400) from exc


def extract_text_from_txt(source) -> str:
    if hasattr(source, "read"):
        data = source.read()
        if isinstance(data, bytes):
            data = data.decode("utf-8", errors="ignore")
        return normalize_text(data)
    return normalize_text(Path(source).read_text(encoding="utf-8", errors="ignore"))


def extract_text(path_or_file, filename: str = "") -> str:
    name = filename or (str(path_or_file) if isinstance(path_or_file, (str, Path)) else "")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if isinstance(path_or_file, (str, Path)):
        handle = open(path_or_file, "rb")
        close = True
    else:
        handle = path_or_file
        close = False
    try:
        if ext == "pdf":
            text = extract_text_from_pdf(handle)
        elif ext in {"doc", "docx"}:
            if hasattr(handle, "seek"):
                handle.seek(0)
            text = extract_text_from_docx(handle)
        elif ext == "txt":
            if hasattr(handle, "seek"):
                handle.seek(0)
            text = extract_text_from_txt(handle)
        else:
            raise APIError("UNSUPPORTED_FORMAT", "Only PDF, DOCX, and TXT files are supported.", 400)
    finally:
        if close:
            handle.close()
    if not text:
        raise APIError("INVALID_DOCUMENT", "No extractable text was found in this file.", 400)
    return text


def _section_blocks(text: str) -> dict[str, str]:
    lines = text.splitlines()
    current = "header"
    buckets: dict[str, list[str]] = {k: [] for k in list(SECTION_HEADERS) + ["header"]}
    for line in lines:
        stripped = line.strip()
        matched = None
        for key, pattern in SECTION_HEADERS.items():
            if pattern.match(stripped):
                matched = key
                break
        if matched:
            current = matched
            continue
        buckets[current].append(stripped)
    return {k: "\n".join(v).strip() for k, v in buckets.items()}


def parse_resume(path_or_file, filename: str = "") -> dict:
    raw = extract_text(path_or_file, filename)
    sections = _section_blocks(raw)
    emails = EMAIL_RE.findall(raw)
    phones = [re.sub(r"\s+", " ", p).strip() for p in PHONE_RE.findall(raw)]
    header_lines = [ln for ln in sections.get("header", "").splitlines() if ln]
    name = header_lines[0] if header_lines else None
    skills = []
    skill_text = sections.get("skills") or ""
    for part in re.split(r"[,•|\n]", skill_text):
        token = part.strip(" -•\t")
        if 1 < len(token) < 48:
            skills.append(token)
    return {
        "name": name,
        "email": emails[0] if emails else None,
        "phone": phones[0] if phones else None,
        "skills": skills[:80],
        "education": sections.get("education") or "",
        "experience": sections.get("experience") or "",
        "projects": sections.get("projects") or "",
        "certifications": sections.get("certifications") or "",
        "raw_text": raw,
    }
