from io import BytesIO

from app.models.job_role import JobRole


def test_summarization_extractive(client, auth_header):
    text = " ".join(
        [
            "Artificial intelligence systems retrieve information from documents.",
            "Summarization compresses long passages while keeping key facts.",
            "Students use roadmaps to sequence skills over several months.",
            "Resume analysis estimates compatibility with a selected job role.",
            "Retrieval augmented generation answers questions from uploaded files.",
        ]
    )
    response = client.post(
        "/api/summarization",
        json={"text": text, "line_count": 3, "method": "extractive"},
        headers=auth_header,
    )
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert data["method"] == "extractive"
    assert data["summary"]


def test_roadmaps_and_resources(client, auth_header):
    roadmaps = client.get("/api/roadmaps", headers=auth_header)
    assert roadmaps.status_code == 200
    items = roadmaps.get_json()["data"]
    assert len(items) >= 5
    detail = client.get(f"/api/roadmaps/{items[0]['id']}", headers=auth_header)
    assert detail.status_code == 200
    resources = client.get("/api/resources", headers=auth_header)
    assert resources.status_code == 200
    assert resources.get_json()["data"]["total"] >= 1


def test_search(client, auth_header):
    response = client.get("/api/search?q=python", headers=auth_header)
    assert response.status_code == 200
    payload = response.get_json()["data"]
    assert "roadmaps" in payload


def test_ats_analyze(client, app, auth_header):
    with app.app_context():
        role = JobRole.query.filter_by(title="Python Developer").first()
        role_id = role.id
    resume = (
        b"Ada Lovelace\nada@example.com\n"
        b"Skills: Python, Flask, SQL, Git, pytest, OOP\n"
        b"Education: Bachelor of Technology in Computer Science\n"
        b"Experience: Software engineer intern building Flask APIs and tests.\n"
    )
    response = client.post(
        "/api/ats/analyze",
        data={"job_role_id": role_id, "file": (BytesIO(resume), "resume.txt")},
        content_type="multipart/form-data",
        headers=auth_header,
    )
    # txt is not allowed for ATS; expect unsupported
    assert response.status_code == 400

    from docx import Document

    buffer = BytesIO()
    doc = Document()
    doc.add_paragraph("Ada Lovelace")
    doc.add_paragraph("ada@example.com")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python, Flask, SQL, Git, pytest, OOP, REST, Authentication, Testing")
    doc.add_paragraph("Education")
    doc.add_paragraph("Bachelor of Technology, Computer Science")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Backend intern building Flask APIs, PostgreSQL models, and pytest suites.")
    doc.save(buffer)
    buffer.seek(0)
    response = client.post(
        "/api/ats/analyze",
        data={"job_role_id": str(role_id), "file": (buffer, "resume.docx")},
        content_type="multipart/form-data",
        headers=auth_header,
    )
    assert response.status_code == 200
    score = response.get_json()["data"]["score"]
    assert 0 <= score <= 100
